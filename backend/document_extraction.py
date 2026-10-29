"""Safe, bounded text extraction for AXION AI knowledge uploads."""
from pathlib import Path
import io
import re

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARS = 6000
SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}

class ExtractionError(ValueError):
    pass


def _clean(text):
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _ocr_image(image):
    try:
        import pytesseract
    except ImportError as exc:
        raise ExtractionError("OCR is not installed. Install backend requirements and the Tesseract OCR application, then retry.") from exc
    try:
        return pytesseract.image_to_string(image)
    except Exception as exc:
        raise ExtractionError("Image OCR failed. Confirm Tesseract OCR is installed and available on PATH.") from exc


def extract_document(filename, raw_bytes, include_metadata=False):
    """Extract text from an allowed document. Optionally include page-level source excerpts."""
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ExtractionError("Unsupported file type. Upload TXT, MD, PDF, DOCX, or a supported image.")
    if not raw_bytes:
        raise ExtractionError("The uploaded file is empty.")
    if len(raw_bytes) > MAX_UPLOAD_BYTES:
        raise ExtractionError("File is larger than the 10 MB upload limit.")

    page_texts = []
    if suffix in {".txt", ".md"}:
        try:
            text = raw_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw_bytes.decode("utf-8", errors="replace")
        method = "text"
    elif suffix == ".docx":
        try:
            from docx import Document
            doc = Document(io.BytesIO(raw_bytes))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            text, method = "\n".join(parts), "docx"
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError("Could not read this Word document. Ensure it is a valid .docx file.") from exc
    elif suffix == ".pdf":
        try:
            import fitz
            pdf = fitz.open(stream=raw_bytes, filetype="pdf")
            parts = [page.get_text("text") for page in pdf]
            page_texts = [{"page_number": index + 1, "text": part} for index, part in enumerate(parts) if part.strip()]
            text = "\n".join(parts)
            method = "pdf_text"
            if len(_clean(text)) < 20:
                try:
                    import pytesseract
                    from PIL import Image
                except ImportError as exc:
                    raise ExtractionError("This PDF appears scanned and needs OCR. Install Pillow, pytesseract, and the Tesseract OCR application.") from exc
                ocr_parts = []
                try:
                    for page in pdf:
                        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                        image = Image.open(io.BytesIO(pix.tobytes("png")))
                        ocr_parts.append(_ocr_image(image))
                except ExtractionError:
                    raise
                except Exception as exc:
                    raise ExtractionError("Could not OCR the scanned PDF. Check that Tesseract OCR is installed and the PDF is valid.") from exc
                page_texts = [{"page_number": index + 1, "text": part} for index, part in enumerate(ocr_parts) if part.strip()]
                text, method = "\n".join(ocr_parts), "pdf_ocr"
            pdf.close()
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError("Could not read this PDF. Ensure it is a valid, unencrypted PDF.") from exc
    else:
        try:
            from PIL import Image, UnidentifiedImageError
            image = Image.open(io.BytesIO(raw_bytes))
            image.verify()
            image = Image.open(io.BytesIO(raw_bytes))
            text, method = _ocr_image(image), "image_ocr"
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError("Could not read this image. Upload a valid image file.") from exc

    text = _clean(text)
    if not text:
        raise ExtractionError("No readable text was found in this document.")
    if len(text) > MAX_EXTRACTED_CHARS:
        text = text[:MAX_EXTRACTED_CHARS].rstrip()
    if include_metadata:
        remaining = MAX_EXTRACTED_CHARS
        bounded_pages = []
        for page in page_texts:
            if remaining <= 0:
                break
            excerpt = _clean(str(page.get("text", "")))[:min(1200, remaining)]
            if excerpt:
                bounded_pages.append({"page_number": page["page_number"], "text": excerpt})
                remaining -= len(excerpt)
        return text, method, bounded_pages
    return text, method
